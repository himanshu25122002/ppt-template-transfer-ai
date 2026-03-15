"""
extractor.py
Extract structured content from PPTX and split large slides into logical slides.
"""

import json
import os
from pptx import Presentation
import re 


MAX_LINES_PER_SLIDE = 6


def extract_lines(text_frame):

    lines = []

    for para in text_frame.paragraphs:

        raw = para.text

        if not raw:
            continue

        # Remove PowerPoint internal vertical tab markers
        clean_text = re.sub(r"_x000B_?", " ", raw)

        # Normalize whitespace
        clean_text = re.sub(r"\s+", " ", clean_text).strip()

        if clean_text:
            lines.append(clean_text)

    return lines


def split_into_slides(title, lines, slide_number_start):
    """Split long content into multiple logical slides."""
    slides = []

    chunk = []
    slide_number = slide_number_start

    for line in lines:
        chunk.append(line)

        if len(chunk) >= MAX_LINES_PER_SLIDE:
            slides.append({
                "slide_number": slide_number,
                "title": title,
                "body": chunk,
                "bullets": chunk
            })
            slide_number += 1
            chunk = []

    if chunk:
        slides.append({
            "slide_number": slide_number,
            "title": title,
            "body": chunk,
            "bullets": chunk
        })

    return slides


def extract_presentation(pptx_path, output_dir="output"):

    os.makedirs(output_dir, exist_ok=True)

    if not os.path.exists(pptx_path):
        raise FileNotFoundError(pptx_path)

    prs = Presentation(pptx_path)

    all_slides = []
    slide_counter = 1

    for slide in prs.slides:

        title = None
        lines = []

        for shape in slide.shapes:

            if not shape.has_text_frame:
                continue

            text_lines = extract_lines(shape.text_frame)

            if not text_lines:
                continue

            if shape == slide.shapes.title and text_lines:
                title = text_lines[0]
            else:
                lines.extend(text_lines)

        if not title and lines:
            title = lines[0]

        logical_slides = split_into_slides(title, lines, slide_counter)

        slide_counter += len(logical_slides)

        all_slides.extend(logical_slides)

    output_path = os.path.join(output_dir, "extracted_content.json")

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(all_slides, f, indent=2, ensure_ascii=False)

    print(f"[Extractor] Extracted {len(all_slides)} logical slides → {output_path}")

    return all_slides


if __name__ == "__main__":

    import sys

    path = sys.argv[1] if len(sys.argv) > 1 else "input/sample.pptx"

    extract_presentation(path)