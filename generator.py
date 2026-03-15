"""
generator.py
Rebuilds a presentation using the selected PPTX template,
applying extracted and classified slide content to matching layouts.
"""

import os
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from copy import deepcopy


# Map classification labels to preferred layout name keywords
LAYOUT_PREFERENCE_MAP = {
    "Title Slide":      ["title slide", "cover", "title, content", "title"],
    "Bullet Slide":     ["title and content", "bullet", "content", "text"],
    "Content Slide":    ["title and content", "content", "text", "blank"],
    "Two Column Slide": ["two content", "comparison", "two column", "content"],
}


def find_best_layout(prs: Presentation, classification: str):
    """
    Find the most appropriate slide layout for the given classification.
    Falls back to the second layout if no keyword match is found.
    """
    keywords = LAYOUT_PREFERENCE_MAP.get(classification, ["title and content"])
    layouts = prs.slide_layouts

    for keyword in keywords:
        for layout in layouts:
            if keyword.lower() in layout.name.lower():
                return layout

    # Safe fallback: layout index 1 is typically "Title and Content"
    return layouts[min(1, len(layouts) - 1)]


def get_placeholder_by_idx(slide, idx: int):
    """Return the placeholder with the given idx, or None."""
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def set_text_safe(placeholder, text: str, font_size: int = None):
    """Set text on a placeholder, clearing existing content first."""
    if placeholder is None:
        return
    tf = placeholder.text_frame
    tf.clear()
    para = tf.paragraphs[0]
    run = para.add_run()
    run.text = text
    if font_size:
        run.font.size = Pt(font_size)


def add_bullet_content(placeholder, bullets: list[str], body_text: list[str]):
    """
    Populate a content placeholder with bullet points and/or body text.
    Bullets take priority; body text is used when bullets are absent.
    """
    if placeholder is None:
        return

    tf = placeholder.text_frame
    tf.clear()
    tf.word_wrap = True

    content_lines = bullets if bullets else body_text

    for i, line in enumerate(content_lines):
        if i == 0:
            para = tf.paragraphs[0]
        else:
            para = tf.add_paragraph()

        # Strip bullet prefix characters added during extraction
        clean_line = line.lstrip("• ").strip()
        run = para.add_run()
        run.text = clean_line

        # Preserve indent level heuristically from leading spaces
        indent = len(line) - len(line.lstrip(" "))
        para.level = min(indent // 2, 4)


def build_title_slide(prs: Presentation, slide_data: dict, layout):
    """Create a title/cover slide."""
    slide = prs.slides.add_slide(layout)

    # Placeholder idx 0 → title, idx 1 → subtitle
    title_ph = get_placeholder_by_idx(slide, 0)
    subtitle_ph = get_placeholder_by_idx(slide, 1)

    title_text = slide_data.get("title") or "Untitled"
    set_text_safe(title_ph, title_text)

    body = slide_data.get("body", [])
    if subtitle_ph and body:
        set_text_safe(subtitle_ph, body[0])

    return slide


def build_content_slide(prs: Presentation, slide_data: dict, layout):
    """Create a standard title + content slide."""
    slide = prs.slides.add_slide(layout)

    title_ph = get_placeholder_by_idx(slide, 0)
    content_ph = get_placeholder_by_idx(slide, 1)

    # Try idx 2 if idx 1 is absent
    if content_ph is None:
        content_ph = get_placeholder_by_idx(slide, 2)

    title_text = slide_data.get("title") or "Slide"
    set_text_safe(title_ph, title_text)

    add_bullet_content(
        content_ph,
        slide_data.get("bullets", []),
        slide_data.get("body", []),
    )

    return slide


def build_two_column_slide(prs: Presentation, slide_data: dict, layout):
    """Create a two-column slide, splitting content across both columns."""
    slide = prs.slides.add_slide(layout)

    title_ph = get_placeholder_by_idx(slide, 0)
    left_ph = get_placeholder_by_idx(slide, 1)
    right_ph = get_placeholder_by_idx(slide, 2)

    set_text_safe(title_ph, slide_data.get("title") or "Slide")

    all_content = slide_data.get("bullets") or slide_data.get("body", [])
    mid = max(1, len(all_content) // 2)

    add_bullet_content(left_ph, all_content[:mid], [])
    add_bullet_content(right_ph, all_content[mid:], [])

    return slide


def generate_presentation(classified_slides, template_path, output_dir="output"):

    os.makedirs(output_dir, exist_ok=True)

    if not os.path.exists(template_path):
        raise FileNotFoundError(template_path)

    # Load template
    prs = Presentation(template_path)

    # SAFE slide removal
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]

    print(f"[Generator] Building {len(classified_slides)} slides using: {os.path.basename(template_path)}")

    for slide_data in classified_slides:

        classification = slide_data.get("classification", "Content Slide")
        layout = find_best_layout(prs, classification)

        print(f"Slide {slide_data['slide_number']} → {classification}")

        if classification == "Title Slide":
            build_title_slide(prs, slide_data, layout)

        elif classification == "Two Column Slide":
            build_two_column_slide(prs, slide_data, layout)

        else:
            build_content_slide(prs, slide_data, layout)

    output_path = os.path.join(output_dir, "final_presentation.pptx")

    prs.save(output_path)

    print(f"[Generator] Saved → {output_path}")

    return output_path


if __name__ == "__main__":
    with open("output/classified_slides.json") as f:
        slides = json.load(f)
    with open("output/selected_template.txt") as f:
        template = f.read().strip()
    generate_presentation(slides, template)
